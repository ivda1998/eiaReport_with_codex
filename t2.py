#!/usr/bin/env python3
import argparse
import datetime as dt
import json
import os
import re
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional

import requests
from docx import Document


ROOT = Path(__file__).resolve().parent
INPUT_DIR = ROOT / "input"
OUTPUT_DIR = ROOT / "output"
DATA_DIR = ROOT / "data"
REFERENCES_DIR = DATA_DIR / "references"
DEFAULT_INPUT = INPUT_DIR / "project.inp"
DEFAULT_OUTPUT = INPUT_DIR / "project.out"
DEFAULT_TEXT_REPORT = OUTPUT_DIR / "대기질_환경영향평가서.txt"
DEFAULT_DOCX_BASENAME = "대기질_환경영향평가서"
PAPERS_JSONL = REFERENCES_DIR / "papers.jsonl"
CITATIONS_BIB = REFERENCES_DIR / "citations.bib"


@dataclass
class PaperEntry:
    title: str
    authors: str
    year: str
    doi: str
    url: str
    source: str
    tags: List[str]
    note: str

    def to_json(self) -> str:
        return json.dumps({
            "title": self.title,
            "authors": self.authors,
            "year": self.year,
            "doi": self.doi,
            "url": self.url,
            "source": self.source,
            "tags": self.tags,
            "note": self.note,
        }, ensure_ascii=False)

    def to_citation_line(self, index: int) -> str:
        author_text = self.authors or "저자미상"
        year_text = self.year or "연도미상"
        title_text = self.title or "제목미상"
        doi_text = f", DOI: {self.doi}" if self.doi else ""
        return f"[{index}] {author_text} ({year_text}). {title_text}.{doi_text}"

    def to_bibtex(self, index: int) -> str:
        key = f"paper{index}"
        fields = {
            "title": self.title,
            "author": self.authors,
            "year": self.year,
            "doi": self.doi,
            "url": self.url,
        }
        body = "\n".join([f"  {k} = {{{v}}}," for k, v in fields.items() if v])
        return f"@article{{{key},\n{body}\n}}"


class PaperLibrary:
    def __init__(self, jsonl_path: Path, bib_path: Path) -> None:
        self.jsonl_path = jsonl_path
        self.bib_path = bib_path
        self.entries: List[PaperEntry] = []

    def load(self) -> None:
        self.entries = []
        if not self.jsonl_path.exists():
            return
        for line in self.jsonl_path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            payload = json.loads(line)
            self.entries.append(
                PaperEntry(
                    title=payload.get("title", ""),
                    authors=payload.get("authors", ""),
                    year=str(payload.get("year", "")),
                    doi=payload.get("doi", ""),
                    url=payload.get("url", ""),
                    source=payload.get("source", ""),
                    tags=payload.get("tags", []),
                    note=payload.get("note", ""),
                )
            )

    def add_entries(self, entries: Iterable[PaperEntry]) -> None:
        self.entries.extend(entries)

    def save(self) -> None:
        if not self.jsonl_path.parent.exists():
            self.jsonl_path.parent.mkdir(parents=True, exist_ok=True)
        lines = [entry.to_json() for entry in self.entries]
        self.jsonl_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        self._save_bibtex()

    def _save_bibtex(self) -> None:
        bib_entries = [entry.to_bibtex(idx + 1) for idx, entry in enumerate(self.entries)]
        self.bib_path.write_text("\n\n".join(bib_entries) + "\n", encoding="utf-8")


def ensure_dirs() -> None:
    for path in [INPUT_DIR, OUTPUT_DIR, DATA_DIR, REFERENCES_DIR]:
        path.mkdir(parents=True, exist_ok=True)


def run_aermod(input_path: Path, output_path: Path, exe_path: Path) -> None:
    if not exe_path.exists():
        return
    output_path.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [str(exe_path), str(input_path)],
        cwd=str(exe_path.parent),
        check=False,
    )


def extract_float_values(text: str) -> List[float]:
    matches = re.findall(r"-?\d+\.\d+", text)
    return [float(value) for value in matches]


def parse_aermod_output(output_path: Path) -> dict:
    if not output_path.exists():
        return {
            "max_concentration": None,
            "sample_count": 0,
        }
    content = output_path.read_text(encoding="utf-8", errors="ignore")
    values = extract_float_values(content)
    return {
        "max_concentration": max(values) if values else None,
        "sample_count": len(values),
    }


def build_report_text(summary: dict, papers: List[PaperEntry]) -> str:
    max_conc = summary.get("max_concentration")
    max_conc_text = f"{max_conc:.4f}" if max_conc is not None else "자료 없음"
    sample_count = summary.get("sample_count", 0)
    lines = [
        "대기질 환경영향평가서",
        "=====================",
        f"작성일: {dt.datetime.now().strftime('%Y-%m-%d')}",
        "",
        "1. 분석 개요",
        "- AERMOD 출력 파일을 기반으로 대기질 영향을 요약했습니다.",
        f"- 추출된 수치 표본 수: {sample_count}",
        "",
        "2. 주요 결과",
        f"- 최대 농도(추정): {max_conc_text}",
        "",
        "3. 참고문헌",
    ]
    if papers:
        for idx, entry in enumerate(papers, start=1):
            lines.append(entry.to_citation_line(idx))
    else:
        lines.append("- 등록된 참고문헌이 없습니다. 학술검색 결과를 추가하세요.")
    return "\n".join(lines)


def write_reports(text_report: Path, docx_basename: str, report_text: str) -> Path:
    text_report.parent.mkdir(parents=True, exist_ok=True)
    text_report.write_text(report_text, encoding="utf-8")
    docx_path = OUTPUT_DIR / f"{docx_basename}_{dt.datetime.now().strftime('%Y%m%d')}.docx"
    document = Document()
    for line in report_text.splitlines():
        if line.strip() == "":
            document.add_paragraph("")
        else:
            document.add_paragraph(line)
    document.save(docx_path)
    return docx_path


def search_crossref(query: str, rows: int) -> List[PaperEntry]:
    response = requests.get(
        "https://api.crossref.org/works",
        params={"query": query, "rows": rows},
        timeout=30,
    )
    response.raise_for_status()
    items = response.json().get("message", {}).get("items", [])
    entries: List[PaperEntry] = []
    for item in items:
        title = " ".join(item.get("title", [])).strip()
        authors_list = item.get("author", [])
        authors = ", ".join(
            [
                " ".join(filter(None, [author.get("family"), author.get("given")]))
                for author in authors_list
            ]
        )
        year = ""
        issued = item.get("issued", {}).get("date-parts", [])
        if issued and issued[0]:
            year = str(issued[0][0])
        doi = item.get("DOI", "")
        url = item.get("URL", "")
        entries.append(
            PaperEntry(
                title=title,
                authors=authors,
                year=year,
                doi=doi,
                url=url,
                source="Crossref",
                tags=["대기질", "AERMOD"],
                note="자동 수집",
            )
        )
    return entries


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="대기질 환경영향평가서 생성 도구")
    subparsers = parser.add_subparsers(dest="command")

    run_parser = subparsers.add_parser("run", help="AERMOD 실행 후 보고서 생성")
    run_parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    run_parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    run_parser.add_argument("--exe", type=Path, default=ROOT / "engine" / "aermod.exe")

    report_parser = subparsers.add_parser("report", help="AERMOD 출력 파일로 보고서 생성")
    report_parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)

    search_parser = subparsers.add_parser("search", help="학술논문 검색 및 저장")
    search_parser.add_argument("query", help="검색 키워드")
    search_parser.add_argument("--rows", type=int, default=5)

    return parser


def main() -> None:
    ensure_dirs()
    parser = build_parser()
    args = parser.parse_args()
    command = args.command or "run"

    library = PaperLibrary(PAPERS_JSONL, CITATIONS_BIB)
    library.load()

    if command == "search":
        entries = search_crossref(args.query, args.rows)
        library.add_entries(entries)
        library.save()
        return

    if command == "run":
        run_aermod(args.input, args.output, args.exe)
        output_path = args.output
    else:
        output_path = args.output

    summary = parse_aermod_output(output_path)
    report_text = build_report_text(summary, library.entries)
    write_reports(DEFAULT_TEXT_REPORT, DEFAULT_DOCX_BASENAME, report_text)


if __name__ == "__main__":
    main()
